import argparse
import datetime
import hashlib
import json
import os
from pathlib import Path
from typing import Dict, Tuple

try:
    from PyPDF2 import PdfReader
    PDF_OK = True
except Exception:
    PdfReader = None
    PDF_OK = False

try:
    from docx import Document
    DOCX_OK = True
except Exception:
    Document = None
    DOCX_OK = False

REQUIRED_SECTIONS = [
    "executive_summary",
    "value_proposition",
    "system_overview",
    "quick_start_guide",
    "template_instructions",
    "examples",
    "license_and_notes",
    "marketing_pack"
]

TOP_LEVEL_KEYS = {
    "schema_version",
    "id",
    "brand",
    "title",
    "slogan",
    "meta_definition",
    "tags",
    "sections",
    "visual_identity",
    "status",
    "version",
    "metadata"
}

METADATA_KEYS = {
    "author",
    "language",
    "source_file",
    "source_name",
    "source_hash",
    "source_bytes",
    "source_mtime",
    "last_updated",
    "pipeline",
    "pipeline_version",
    "run_id",
    "pii_redacted",
    "anonymized_source"
}

SCRIPT_DIR = Path(__file__).resolve().parent
REPO_ROOT = SCRIPT_DIR.parent

def hash_text(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8", errors="ignore")).hexdigest()

def read_text_file(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="ignore").strip()

def read_pdf(path: Path) -> str:
    if not PDF_OK:
        return ""
    try:
        reader = PdfReader(str(path))
        text = "\n".join((page.extract_text() or "") for page in reader.pages)
        return text.strip()
    except Exception:
        return ""

def read_docx(path: Path) -> str:
    if not DOCX_OK:
        return ""
    try:
        doc = Document(str(path))
        return "\n".join([para.text for para in doc.paragraphs]).strip()
    except Exception:
        return ""

def read_source_file(path: Path) -> str:
    ext = path.suffix.lower()
    if ext in (".txt", ".md"):
        return read_text_file(path)
    if ext == ".pdf":
        return read_pdf(path)
    if ext == ".docx":
        return read_docx(path)
    return ""

def coerce_sections(sections: object, placeholder: str) -> Tuple[Dict[str, str], int]:
    placeholder_count = 0
    if not isinstance(sections, dict):
        sections = {}
    normalized: Dict[str, str] = {}
    for key in REQUIRED_SECTIONS:
        value = sections.get(key)
        if isinstance(value, str) and value.strip():
            normalized[key] = value
        else:
            normalized[key] = placeholder
            placeholder_count += 1
    return normalized, placeholder_count

def resolve_source_path(value: str) -> Path:
    source_path = Path(value)
    if source_path.is_absolute():
        return source_path
    return REPO_ROOT / source_path

def fallback_source_text(data: dict) -> str:
    if isinstance(data.get("content"), str):
        return data["content"]
    sections = data.get("sections") or {}
    if isinstance(sections, dict):
        summary = sections.get("executive_summary")
        if isinstance(summary, str):
            return summary
        return json.dumps(sections, ensure_ascii=False)
    return json.dumps(data, ensure_ascii=False)

def migrate_blueprint(
    data: dict,
    path: Path,
    schema_version: str,
    pipeline_version: str,
    placeholder: str,
    use_source: bool
) -> Tuple[dict, Dict[str, int]]:
    report = {
        "placeholders_added": 0,
        "source_read": 0,
        "fields_added": 0
    }

    sections, placeholders = coerce_sections(data.get("sections"), placeholder)
    report["placeholders_added"] += placeholders

    metadata = data.get("metadata")
    if not isinstance(metadata, dict):
        metadata = {}

    source_file = metadata.get("source_file") or f"blueprints/{path.name}"
    source_file = str(source_file)
    source_name = metadata.get("source_name")
    if not source_name:
        try:
            source_name = Path(source_file).name
        except Exception:
            source_name = path.name

    source_text = ""
    if use_source and source_file:
        source_path = resolve_source_path(source_file)
        if source_path.exists():
            source_text = read_source_file(source_path)
            if source_text:
                report["source_read"] += 1
                if "source_mtime" not in metadata:
                    metadata["source_mtime"] = datetime.datetime.fromtimestamp(
                        source_path.stat().st_mtime
                    ).isoformat()

    if not source_text:
        source_text = fallback_source_text(data)

    source_hash = metadata.get("source_hash") or hash_text(source_text)
    source_bytes = metadata.get("source_bytes")
    if not isinstance(source_bytes, int):
        source_bytes = len(source_text.encode("utf-8", errors="ignore"))

    author = metadata.get("author") or "Unknown"
    language = metadata.get("language") or "en"
    pipeline = metadata.get("pipeline") or "unknown"
    last_updated = metadata.get("last_updated") or str(datetime.date.today())
    pipeline_version_value = metadata.get("pipeline_version") or pipeline_version

    new_metadata = {
        "author": author,
        "language": language,
        "source_file": source_file,
        "source_name": source_name,
        "source_hash": source_hash,
        "source_bytes": source_bytes,
        "last_updated": last_updated,
        "pipeline": pipeline,
        "pipeline_version": pipeline_version_value
    }

    for key in ("source_mtime", "run_id", "pii_redacted", "anonymized_source"):
        if key in metadata:
            new_metadata[key] = metadata[key]

    filtered = {
        "schema_version": schema_version,
        "id": data.get("id") or f"BP-{hash_text(path.name)[:8]}",
        "brand": data.get("brand") or "NamoNexus",
        "title": data.get("title") or path.stem,
        "slogan": data.get("slogan") or "",
        "meta_definition": data.get("meta_definition") or "",
        "tags": data.get("tags") if isinstance(data.get("tags"), list) else [],
        "sections": sections,
        "visual_identity": data.get("visual_identity") if isinstance(data.get("visual_identity"), dict) else {},
        "status": data.get("status") or "draft",
        "version": data.get("version") or "0.1",
        "metadata": new_metadata
    }

    report["fields_added"] = len([k for k in TOP_LEVEL_KEYS if k not in data])
    return filtered, report

def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Migrate blueprint JSON files to the latest schema.")
    parser.add_argument("--input-dir", default=str(REPO_ROOT / "blueprints"), help="Blueprints directory.")
    parser.add_argument("--schema-version", default="1.1", help="Schema version to apply.")
    parser.add_argument("--pipeline-version", default="legacy", help="Pipeline version to apply when missing.")
    parser.add_argument("--placeholder", default="Placeholder - missing in source", help="Placeholder for missing sections.")
    parser.add_argument("--max-files", type=int, default=0, help="Limit number of files processed.")
    parser.add_argument("--apply", action="store_true", help="Write changes to disk.")
    parser.add_argument("--use-source", action="store_true", help="Attempt to read source files for hashes.")
    parser.add_argument("--report", default="", help="Write a JSON report to this path.")
    return parser.parse_args()

def main() -> int:
    args = parse_args()
    input_dir = Path(args.input_dir)
    if not input_dir.exists():
        print(f"Input directory not found: {input_dir}")
        return 1

    files = sorted([p for p in input_dir.iterdir() if p.suffix.lower() == ".json"])
    if args.max_files and args.max_files > 0:
        files = files[:args.max_files]

    summary = {
        "total": len(files),
        "updated": 0,
        "unchanged": 0,
        "errors": 0,
        "placeholders_added": 0,
        "source_read": 0
    }

    for path in files:
        try:
            original = json.loads(path.read_text(encoding="utf-8"))
        except Exception:
            summary["errors"] += 1
            continue

        migrated, report = migrate_blueprint(
            original,
            path,
            args.schema_version,
            args.pipeline_version,
            args.placeholder,
            args.use_source
        )
        summary["placeholders_added"] += report["placeholders_added"]
        summary["source_read"] += report["source_read"]

        if migrated == original:
            summary["unchanged"] += 1
            continue

        summary["updated"] += 1
        if args.apply:
            path.write_text(json.dumps(migrated, ensure_ascii=False, indent=2), encoding="utf-8")

    print(
        "Migration summary:",
        f"total={summary['total']}",
        f"updated={summary['updated']}",
        f"unchanged={summary['unchanged']}",
        f"errors={summary['errors']}",
        f"placeholders={summary['placeholders_added']}",
        f"source_read={summary['source_read']}"
    )

    if args.report:
        Path(args.report).write_text(json.dumps(summary, ensure_ascii=True, indent=2), encoding="utf-8")

    return 0

if __name__ == "__main__":
    raise SystemExit(main())
