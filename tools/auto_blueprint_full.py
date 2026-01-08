import argparse
import datetime
import hashlib
import json
import logging
import os
import re
import time
import uuid
from pathlib import Path
from typing import Dict, List, Tuple, Optional

try:
    from google import genai
    GENAI_OK = True
except Exception:
    genai = None
    GENAI_OK = False

try:
    from PyPDF2 import PdfReader
    PDF_OK = True
except Exception:
    PdfReader = None
    PDF_OK = False

try:
    from docx import Document  # python-docx
    DOCX_OK = True
except Exception:
    Document = None
    DOCX_OK = False

from sanitize import sanitize_text
from pii import redact_pii

# -------- Config --------
SCRIPT_DIR = Path(__file__).resolve().parent
REPO_ROOT = SCRIPT_DIR.parent

DEFAULT_INPUT_DIR = REPO_ROOT / "framework"
DEFAULT_OUTPUT_DIR = REPO_ROOT / "blueprints"
DEFAULT_SCHEMA_PATH = SCRIPT_DIR / "schema_blueprint.json"

PIPELINE_NAME = "auto-blueprint-full"
PIPELINE_VERSION = "1.1"
SCHEMA_VERSION = "1.1"

BRAND = "NamoNexus"
SLOGAN = "Elevate your existence with NamoNexus."
META_DEF = (
    "This Blueprint is designed as an entity beyond AI - "
    "a self-evolving, meta-intelligent framework that grows infinitely across dimensions."
)

TIMEOUT = 20
RETRIES = 3

# -------- Gemini Setup --------
def configure_gemini(enable_llm: bool):
    if not enable_llm:
        logging.info("LLM enrichment disabled; skipping Gemini.")
        return None
    api_key = os.getenv("GEMINI_API_KEY") or os.getenv("GEMENI_KEY")
    if api_key and GENAI_OK:
        return genai.Client(api_key=api_key)
    if api_key and not GENAI_OK:
        logging.warning(
            "GEMINI_API_KEY set but google-genai is not installed; skipping Gemini enrichment."
        )
        return None
    logging.warning("GEMINI_API_KEY not found; skipping Gemini enrichment.")
    return None

# -------- Readers --------
def read_txt_md(p: Path) -> str:
    return p.read_text(encoding="utf-8", errors="ignore").strip()

def read_pdf(p: Path) -> str:
    if not PDF_OK:
        logging.warning("PyPDF2 not installed; skip .pdf")
        return ""
    try:
        reader = PdfReader(str(p))
        text = "\n".join((page.extract_text() or "") for page in reader.pages)
        return text.strip()
    except Exception as e:
        logging.warning(f"PDF read error {p}: {e}")
        return ""

def read_docx(p: Path) -> str:
    if not DOCX_OK:
        logging.warning("python-docx not installed; skip .docx")
        return ""
    try:
        d = Document(str(p))
        return "\n".join([para.text for para in d.paragraphs]).strip()
    except Exception as e:
        logging.warning(f"DOCX read error {p}: {e}")
        return ""

def load_raw(p: Path) -> str:
    ext = p.suffix.lower()
    if ext in [".txt", ".md"]:
        return read_txt_md(p)
    if ext == ".pdf":
        return read_pdf(p)
    if ext == ".docx":
        return read_docx(p)
    return ""

# -------- Helpers --------
def hash_text(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8", errors="ignore")).hexdigest()

def hash_identifier(value: str) -> str:
    return hashlib.sha256(value.encode("utf-8", errors="ignore")).hexdigest()[:16]

def relative_source_path(src: Path, repo_root: Path) -> str:
    try:
        rel = src.relative_to(repo_root)
        return rel.as_posix()
    except Exception:
        return src.as_posix()

def sanitize_filename(filename: str) -> str:
    """
    Sanitizes a filename by forcing it to ASCII, removing special characters,
    and handling potential empty results.
    """
    stem = Path(filename).stem
    ascii_stem = stem.encode("ascii", "ignore").decode("ascii")
    safe_stem = re.sub(r"[^\w.-]+", "_", ascii_stem)
    safe_stem = re.sub(r"_+", "_", safe_stem).strip("_")
    if not safe_stem:
        return f"blueprint_{hashlib.md5(stem.encode()).hexdigest()[:8]}"
    return safe_stem

def build_output_path(output_dir: Path, safe_stem: str, source_hash: str, suffix: str = "") -> Path:
    base_name = f"{safe_stem}{suffix}"
    primary = output_dir / f"{base_name}.json"
    if not primary.exists():
        return primary
    existing_hash = extract_source_hash(primary)
    if existing_hash == source_hash:
        return primary
    suffix = source_hash[:8]
    candidate = output_dir / f"{base_name}_{suffix}.json"
    if not candidate.exists():
        return candidate
    counter = 2
    while True:
        candidate = output_dir / f"{base_name}_{suffix}_{counter}.json"
        if not candidate.exists():
            return candidate
        counter += 1

def extract_source_hash(path: Path) -> str:
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return ""
    metadata = data.get("metadata", {})
    return metadata.get("segment_hash", "") or metadata.get("source_hash", "")

def format_mtime(src: Path) -> str:
    try:
        return datetime.datetime.fromtimestamp(src.stat().st_mtime).isoformat()
    except Exception:
        return ""

# -------- Blueprint Maker --------
def make_id(seed: str) -> str:
    h = hashlib.sha1(seed.encode("utf-8", errors="ignore")).hexdigest()[:6]
    return f"BP-{datetime.date.today().strftime('%Y%m%d')}-{h}"

def scaffold_sections(content: str) -> dict:
    return {
        "executive_summary": content[:600],
        "value_proposition": "Universal adaptability, modular architecture, AI-agnostic integration, and self-evolving design.",
        "system_overview": "A layered framework transforming raw data into commercial blueprints with metadata and validation.",
        "quick_start_guide": "1) Place raw files into /framework  2) Run the pipeline  3) Review JSON in /blueprints  4) Deploy to your ecosystem.",
        "template_instructions": "Follow the schema. Keep sections concise. Use neutral, professional English. Avoid personal names.",
        "examples": "Education curricula, healthcare protocols, financial decision flows, brand guidelines.",
        "license_and_notes": "Licensed under NamoVerse Creative Framework License (NCFL-1.0). Provide attribution for redistribution.",
        "marketing_pack": "Target: builders, strategists, educators. Pain: unstructured data. USP: chaos-to-commerce via meta-intelligence. Pricing: Base/Pro tiers. GTM: ProductHunt + LinkedIn."
    }

def build_source_info(src: Path, raw_text: str, anonymize_source: bool) -> Dict[str, object]:
    source_file = relative_source_path(src, REPO_ROOT)
    source_name = src.name
    source_hash = hash_text(raw_text)
    source_bytes = len(raw_text.encode("utf-8", errors="ignore"))
    source_mtime = format_mtime(src)
    if anonymize_source:
        source_id = hash_identifier(source_file)
        source_file = f"hash:{source_id}"
        source_name = f"hash:{source_id}"
    return {
        "source_file": source_file,
        "source_name": source_name,
        "source_hash": source_hash,
        "source_bytes": source_bytes,
        "source_mtime": source_mtime,
        "anonymized_source": anonymize_source
    }

def build_blueprint(
    src: Path,
    content: str,
    source_info: Dict[str, object],
    run_id: str,
    pii_redacted: bool,
    segment_index: int,
    segment_count: int,
    segment_title: Optional[str],
    segment_hash: str,
    content_type: str,
    role: str = "",
    sections_override: Optional[Dict[str, str]] = None,
    extra_tags: Optional[List[str]] = None,
    license_id: str = "",
    build_id: str = ""
) -> dict:
    base_title = src.stem.replace("_", " ").strip()
    title = segment_title or base_title
    metadata = {
        "author": "NamoVerse Engine",
        "language": "en",
        "source_file": source_info["source_file"],
        "source_name": source_info["source_name"],
        "source_hash": source_info["source_hash"],
        "source_bytes": source_info["source_bytes"],
        "source_mtime": source_info["source_mtime"],
        "last_updated": str(datetime.date.today()),
        "pipeline": PIPELINE_NAME,
        "pipeline_version": PIPELINE_VERSION,
        "pii_redacted": pii_redacted,
        "anonymized_source": source_info["anonymized_source"],
        "segment_index": segment_index,
        "segment_count": segment_count,
        "segment_title": segment_title or "",
        "segment_hash": segment_hash,
        "content_type": content_type,
        "role": role,
        "license_id": license_id,
        "build_id": build_id
    }
    if run_id:
        metadata["run_id"] = run_id

    return {
        "schema_version": SCHEMA_VERSION,
        "id": make_id(f"{src.name}:{segment_index}:{segment_hash}"),
        "brand": BRAND,
        "title": title,
        "slogan": SLOGAN,
        "meta_definition": META_DEF,
        "tags": [t for t in [content_type, role] if t] + (extra_tags or []),
        "sections": sections_override or scaffold_sections(content),
        "visual_identity": {},
        "status": "draft",
        "version": "0.2",
        "metadata": metadata
    }

def extract_openapi_metadata(text: str) -> Dict[str, str]:
    title_match = re.search(r"^\s*title:\s*(.+)$", text, re.MULTILINE)
    version_match = re.search(r"^\s*version:\s*(.+)$", text, re.MULTILINE)
    desc_match = re.search(r"^\s*description:\s*\|?\s*$\n([\s\S]{0,800})", text, re.MULTILINE)
    title = (title_match.group(1).strip() if title_match else "API Blueprint")
    version = (version_match.group(1).strip() if version_match else "0.1.0")
    description = ""
    if desc_match:
        description = desc_match.group(1).strip().replace("\n", " ")
        description = re.sub(r"\s+", " ", description)
    return {"title": title, "version": version, "description": description}

def scaffold_master_blueprint(content: str) -> Dict[str, str]:
    meta = extract_openapi_metadata(content)
    summary = meta["description"] or content[:600]
    return {
        "executive_summary": (
            f"{meta['title']} defines a commercial-grade memory API blueprint that teams can adopt "
            f"to ship memory features faster and with lower integration risk. {summary} "
            "It standardizes how memory is stored, retrieved, and governed so product teams can "
            "monetize and scale knowledge features without rewriting contracts per project."
        ),
        "value_proposition": (
            "Production-ready API blueprint that reduces ambiguity, compresses development time, "
            "and creates a repeatable memory product offering."
        ),
        "system_overview": (
            "Layered API design with clear contracts for storage, retrieval, decay, and governance. "
            "The blueprint aligns product, engineering, and compliance requirements in a single spec."
        ),
        "quick_start_guide": (
            "1) Review API scope and required endpoints  2) Map endpoints to your stack  "
            "3) Implement adapters and storage backends  4) Validate with test payloads  "
            "5) Pilot with one service before platform rollout."
        ),
        "template_instructions": (
            "Keep endpoints stable. Add examples and error codes. Track versioning explicitly. "
            "Document auth, rate limits, and data retention policy."
        ),
        "examples": (
            "Memory upsert, semantic retrieval, decay scheduling, audit logging, "
            "feedback loops for continuous improvement."
        ),
        "license_and_notes": (
            "Licensed under NamoVerse Creative Framework License (NCFL-1.0). Provide attribution "
            "for redistribution. Include watermark metadata (license_id, build_id) in releases."
        ),
        "marketing_pack": (
            "Target: AI platform teams and integrators. Pain: inconsistent memory contracts and "
            "slow delivery. USP: unified memory API blueprint with governance baked in. "
            "Pricing: setup + support. GTM: pilot-first, then scale across services."
        ),
        "engineering_alignment": (
            "Modules: Memory Core, Retrieval Engine, Decay Scheduler, Audit Logger. "
            "Interfaces: REST/OpenAPI + event hooks. Required: schema definitions, error model, "
            "and compatibility tests."
        ),
        "self_evolution_loop": (
            "Events -> feature extraction -> memory upsert -> retrieval -> response generation -> "
            "evaluation -> policy update -> canary deploy."
        ),
        "kpis": (
            "Latency p95, error rate, retrieval MRR@k, conflict rate, and tone suitability."
        ),
        "safety_compliance": (
            "PII redaction, rate limits, audit logs, and access controls for sensitive memory data."
        )
    }

def scaffold_technical_spec(content: str) -> Dict[str, str]:
    meta = extract_openapi_metadata(content)
    summary = meta["description"] or "OpenAPI-based contract for memory services."
    return {
        "executive_summary": (
            f"{meta['title']} v{meta['version']} defines the contract for memory operations and "
            "enforces consistent request/response behavior across services."
        ),
        "value_proposition": "Clear interface boundaries for memory storage, retrieval, and policy enforcement.",
        "system_overview": (
            f"{summary} Use as a reference contract across services and deployment environments. "
            "Ensure schema consistency and explicit versioning."
        ),
        "quick_start_guide": (
            "1) Import OpenAPI spec  2) Generate client/server stubs  "
            "3) Implement storage adapters  4) Run contract tests  5) Add observability hooks."
        ),
        "template_instructions": (
            "Document request/response bodies, error codes, auth requirements, and pagination. "
            "Provide deterministic examples and edge cases."
        ),
        "examples": "POST /memory/upsert, GET /memory/retrieve, POST /memory/decay",
        "license_and_notes": (
            "Include audit logging, rate limits, and PII redaction policy. "
            "Track schema changes via semantic versioning."
        ),
        "marketing_pack": "Internal spec for engineering alignment and delivery.",
        "self_evolution_loop": (
            "Collect usage events -> analyze retrieval quality -> adjust weights -> "
            "test in canary -> deploy when KPI improves."
        ),
        "kpis": (
            "Latency p95 < 400ms, error rate < 1%, retrieval MRR@k, conflict rate, "
            "schema drift alerts."
        ),
        "safety_compliance": (
            "PII redaction, access controls, audit trails, and rate limiting."
        )
    }

def scaffold_marketing_one_pager(content: str) -> Dict[str, str]:
    meta = extract_openapi_metadata(content)
    return {
        "executive_summary": (
            f"{meta['title']} is a deployable memory API blueprint that reduces time-to-integration "
            "and lets teams productize memory features quickly."
        ),
        "value_proposition": (
            "Ship memory features faster with a proven contract, clear integration steps, and "
            "governance-ready defaults."
        ),
        "system_overview": (
            "Developer-first API spec with clear endpoints, versioning, and integration guidance "
            "for platform rollouts."
        ),
        "quick_start_guide": "1) Evaluate fit  2) Pilot in one service  3) Expand to platform.",
        "template_instructions": "Keep language simple. Highlight outcomes and integration speed.",
        "examples": "AI copilots, knowledge assistants, retrieval pipelines.",
        "license_and_notes": "Commercial-ready blueprint with governance guidance and compliance notes.",
        "marketing_pack": (
            "Target: platform leaders and solution architects. USP: faster rollout and consistent "
            "memory contracts. GTM: pilot-first with measurable ROI."
        )
    }

def compute_code_score(text: str) -> float:
    lines = [line for line in text.splitlines() if line.strip()]
    if not lines:
        return 0.0
    code_hits = 0
    code_patterns = [
        r"^\s*(def|class|import|from|#include|using|public|private|protected|function|const|var|let)\b",
        r"[{};]",
        r"^\s*<\w+[^>]*>",
        r"^\s*[\w\-]+:\s*.+"
    ]
    for line in lines:
        if any(re.search(pat, line) for pat in code_patterns):
            code_hits += 1
    return code_hits / max(len(lines), 1)

def classify_content(text: str) -> str:
    sample = text.lower()
    code_score = compute_code_score(text)
    if code_score >= 0.25:
        return "code"
    if any(k in sample for k in ["self-evolution", "evolution", "ab test", "canary", "kpi", "drift", "metrics"]):
        return "evolution"
    if any(k in sample for k in ["prompt", "system message", "role", "instruction", "template"]):
        return "prompt"
    if any(k in sample for k in ["architecture", "module", "interface", "api", "schema"]):
        return "architecture"
    return "blueprint"

def is_set_header(line: str) -> bool:
    line = line.strip()
    if not line or len(line) > 140:
        return False
    patterns = [
        r"^(ชุดที่|Set|Part|Module|Phase|Section|บทที่)\s*([0-9]+|[IVX]+)\b",
        r"^\d{1,2}[.)]\s+\S+",
        r"^#{1,3}\s+\S+"
    ]
    return any(re.match(pat, line, re.IGNORECASE) for pat in patterns)

def split_content_sets(text: str) -> List[Dict[str, str]]:
    if not text.strip():
        return []
    code_score = compute_code_score(text)
    strict_headers = code_score >= 0.25
    lines = text.splitlines()
    segments: List[Dict[str, str]] = []
    current_lines: List[str] = []
    current_title: Optional[str] = None

    for line in lines:
        header_hit = is_set_header(line)
        if header_hit and strict_headers and not re.match(r"^#{1,3}\s+\S+", line.strip()):
            header_hit = False
        if header_hit and current_lines:
            segments.append({
                "title": (current_title or "").strip(),
                "content": "\n".join(current_lines).strip()
            })
            current_lines = [line]
            current_title = line.strip()
            continue
        if header_hit and not current_lines:
            current_title = line.strip()
            current_lines.append(line)
            continue
        current_lines.append(line)

    if current_lines:
        segments.append({
            "title": (current_title or "").strip(),
            "content": "\n".join(current_lines).strip()
        })

    if len(segments) <= 1:
        return [{"title": "", "content": text.strip()}]
    return segments

def resolve_output_dir(base_dir: Path, content_type: str, layout: str) -> Path:
    if layout == "by-type":
        target = base_dir / content_type
        target.mkdir(parents=True, exist_ok=True)
        return target
    return base_dir

# -------- Gemini Integration --------
def gemini_enrich(blueprint: dict, client) -> dict:
    if client is None or not GENAI_OK:
        if client is None:
            logging.info("Gemini client not configured; skipping enrichment.")
        return blueprint

    raw_content = blueprint["sections"]["executive_summary"]
    title = blueprint["title"]

    prompt = f"""
You are an expert Knowledge Architect. Your task is to transform raw unstructured text into a structured "Blueprint" JSON format.

Here is the raw content from a document titled "{title}":
---
{raw_content[:4000]}
---
(Note: Content truncated to first 4000 characters if too long)

Please analyze the content and generate a JSON object that strictly follows this structure (do not include markdown fencing, just the JSON):
{{
  "sections": {{
    "executive_summary": "A concise summary of the core idea (max 3 sentences).",
    "value_proposition": "What is the unique value or benefit? (max 2 sentences).",
    "system_overview": "Technical or logical architecture description.",
    "quick_start_guide": "Step-by-step guide to get started.",
    "examples": "Practical use cases.",
    "marketing_pack": "Target audience, pain points, and selling points."
  }},
  "tags": ["tag1", "tag2", "tag3"]
}}

If the raw content is empty or meaningless, return reasonable placeholders related to "{title}".
"""

    def list_supported_models() -> List[str]:
        try:
            models = client.models.list()
            names = []
            for m in models:
                name = getattr(m, "name", "") or ""
                supported = getattr(m, "supported_generation_methods", []) or []
                if "generateContent" in supported and name:
                    names.append(name.replace("models/", ""))
            return names
        except Exception as e:
            logging.warning(f"Model list failed: {e}")
            return []

    model_override = os.getenv("GEMINI_MODEL", "").strip()
    model_candidates = [model_override] if model_override else []
    model_candidates += list_supported_models()
    if not model_candidates:
        model_candidates = [
            "gemini-2.0-flash",
            "gemini-2.0-flash-lite",
            "gemini-1.5-flash",
            "gemini-1.5-pro"
        ]

    last_error = None
    for model_name in model_candidates:
        for attempt in range(1, RETRIES + 1):
            try:
                response = client.models.generate_content(
                    model=model_name,
                    contents=prompt
                )
                text = (response.text or "").strip()
                if text.startswith("```json"):
                    text = text[7:-3]
                if text.startswith("```"):
                    text = text[3:-3]

                data = json.loads(text)
                blueprint["sections"].update(data.get("sections", {}))
                blueprint["tags"] = data.get("tags", []) or []
                logging.info(f"Gemini enrichment success using {model_name}.")
                return blueprint
            except Exception as e:
                last_error = e
                logging.warning(f"Gemini error ({model_name}) attempt {attempt}: {e}")
                time.sleep(2 * attempt)

    logging.warning(f"Gemini enrichment failed after retries; using base blueprint. Last error: {last_error}")
    return blueprint

# -------- Runner --------
def process_one(
    p: Path,
    output_dir: Path,
    client,
    run_id: str,
    redact: bool,
    anonymize_source: bool,
    dry_run: bool,
    skip_unchanged: bool,
    max_bytes: int,
    output_lock,
    output_layout: str,
    output_mode: str,
    output_format: str,
    license_id: str,
    build_id: str
) -> Tuple[Dict[str, object], bool]:
    started = time.time()
    result: Dict[str, object] = {
        "source_file": "",
        "source_name": "",
        "source_hash": "",
        "segment_count": 0,
        "status": "error",
        "output_file": "",
        "warnings": [],
        "errors": [],
        "duration_ms": 0
    }

    if max_bytes and p.stat().st_size > max_bytes:
        result["status"] = "skipped"
        result["warnings"].append("max-bytes-exceeded")
        result["duration_ms"] = int((time.time() - started) * 1000)
        return result, True

    raw_text = load_raw(p)
    if not raw_text:
        result["status"] = "skipped"
        result["warnings"].append("empty-or-unreadable")
        result["duration_ms"] = int((time.time() - started) * 1000)
        return result, True

    sanitized = sanitize_text(raw_text)
    pii_redacted = False
    if redact:
        sanitized, pii_stats = redact_pii(sanitized)
        pii_redacted = any(pii_stats.values())

    source_info = build_source_info(p, raw_text, anonymize_source)
    result["source_file"] = source_info["source_file"]
    result["source_name"] = source_info["source_name"]
    result["source_hash"] = source_info["source_hash"]

    safe_stem = sanitize_filename(p.name)

    segments = split_content_sets(sanitized)
    if not segments:
        result["status"] = "skipped"
        result["warnings"].append("empty-segments")
        result["duration_ms"] = int((time.time() - started) * 1000)
        return result, True

    result["segment_count"] = len(segments)
    output_files: List[str] = []

    for idx, seg in enumerate(segments, start=1):
        segment_text = seg["content"]
        if not segment_text:
            continue
        segment_hash = hash_text(segment_text)
        segment_title = seg["title"] or ""
        content_type = classify_content(segment_text)
        target_dir = resolve_output_dir(output_dir, content_type, output_layout)
        suffix = ""
        if len(segments) > 1:
            suffix = f"_set{idx:02d}"
        role_sets: List[Tuple[str, Dict[str, str]]] = []
        if output_mode == "role-split":
            role_sets = [
                ("master", scaffold_master_blueprint(segment_text)),
                ("technical", scaffold_technical_spec(segment_text)),
                ("marketing", scaffold_marketing_one_pager(segment_text))
            ]
        else:
            role_sets = [("", {})]

        for role_name, sections in role_sets:
            role_suffix = f"_{role_name}" if role_name else ""
            out_path = build_output_path(target_dir, safe_stem, segment_hash, suffix + role_suffix)
            if skip_unchanged and out_path.exists() and output_format in ("json", "both"):
                existing_hash = extract_source_hash(out_path)
                if existing_hash == segment_hash:
                    output_files.append(out_path.as_posix())
                    continue

            bp = build_blueprint(
                p,
                segment_text,
                source_info,
                run_id,
                pii_redacted,
                idx,
                len(segments),
                segment_title,
                segment_hash,
                content_type,
                role=role_name,
                sections_override=sections or None,
                extra_tags=["role-split"] if role_name else [],
                license_id=license_id,
                build_id=build_id
            )
            bp = gemini_enrich(bp, client)
            bp["status"] = "complete"

            if dry_run:
                output_files.append(out_path.as_posix())
                continue

            try:
                if output_lock:
                    output_lock.acquire()
                if output_format in ("json", "both"):
                    out_path = build_output_path(target_dir, safe_stem, segment_hash, suffix + role_suffix)
                    out_path.write_text(json.dumps(bp, ensure_ascii=False, indent=2), encoding="utf-8")
                    output_files.append(out_path.as_posix())
                    logging.info(f"Saved blueprint: {out_path}")
                if output_format in ("md", "both"):
                    md_path = build_output_path(target_dir, safe_stem, segment_hash, suffix + role_suffix).with_suffix(".md")
                    md_path.write_text(blueprint_to_markdown(bp), encoding="utf-8")
                    output_files.append(md_path.as_posix())
                    logging.info(f"Saved blueprint: {md_path}")
            except Exception as e:
                result["status"] = "error"
                result["errors"].append(f"write-error:{e}")
                result["duration_ms"] = int((time.time() - started) * 1000)
                return result, False
            finally:
                if output_lock:
                    output_lock.release()

    if dry_run:
        result["status"] = "dry_run"
    else:
        result["status"] = "ok"
    result["output_file"] = ";".join(output_files)
    result["duration_ms"] = int((time.time() - started) * 1000)
    return result, True

def write_manifest(path: Path, payload: Dict[str, object]) -> None:
    path.write_text(json.dumps(payload, ensure_ascii=True, indent=2), encoding="utf-8")

def write_audit_log(path: Path, entries: List[Dict[str, object]]) -> None:
    with path.open("a", encoding="utf-8") as handle:
        for entry in entries:
            handle.write(json.dumps(entry, ensure_ascii=True) + "\n")

def blueprint_to_markdown(blueprint: Dict[str, object]) -> str:
    sections = blueprint.get("sections", {})
    title = blueprint.get("title", "Blueprint")
    brand = blueprint.get("brand", "")
    slogan = blueprint.get("slogan", "")
    meta_def = blueprint.get("meta_definition", "")
    lines = [
        f"# {title}",
        "",
        f"Brand: {brand}" if brand else "",
        f"Slogan: {slogan}" if slogan else "",
        "",
        "Meta Definition:",
        meta_def,
        ""
    ]
    for key, label in [
        ("executive_summary", "Executive Summary"),
        ("value_proposition", "Value Proposition"),
        ("system_overview", "System Overview"),
        ("quick_start_guide", "Quick Start Guide"),
        ("template_instructions", "Template Instructions"),
        ("examples", "Examples"),
        ("license_and_notes", "License and Notes"),
        ("marketing_pack", "Marketing Pack"),
        ("engineering_alignment", "Engineering Alignment"),
        ("self_evolution_loop", "Self-Evolution Loop"),
        ("kpis", "KPIs"),
        ("safety_compliance", "Safety and Compliance")
    ]:
        content = sections.get(key, "")
        if content:
            lines.extend([f"## {label}", content, ""])
    return "\n".join([line for line in lines if line is not None])

def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Auto-generate blueprint JSON files from documents.")
    parser.add_argument("--input-dir", default=str(DEFAULT_INPUT_DIR), help="Input directory with raw files.")
    parser.add_argument("--output-dir", default=str(DEFAULT_OUTPUT_DIR), help="Output directory for blueprints.")
    parser.add_argument("--schema", default=str(DEFAULT_SCHEMA_PATH), help="Schema path (informational).")
    parser.add_argument("--manifest", default="", help="Path to manifest JSON output.")
    parser.add_argument("--audit-log", default="", help="Path to audit log JSONL.")
    parser.add_argument("--dry-run", action="store_true", help="Process files without writing outputs.")
    parser.add_argument("--max-files", type=int, default=0, help="Limit number of files processed.")
    parser.add_argument("--max-bytes", type=int, default=0, help="Skip files larger than this size.")
    parser.add_argument("--skip-unchanged", action="store_true", help="Skip files if output hash matches.")
    parser.add_argument(
        "--redact-pii",
        action=argparse.BooleanOptionalAction,
        default=True,
        help="Redact obvious PII before output (default: true)."
    )
    parser.add_argument(
        "--anonymize-source",
        action=argparse.BooleanOptionalAction,
        default=True,
        help="Hash source paths in metadata (default: true)."
    )
    parser.add_argument(
        "--enable-llm",
        action=argparse.BooleanOptionalAction,
        default=False,
        help="Enable LLM enrichment (default: false)."
    )
    parser.add_argument(
        "--license-id",
        default=os.getenv("LICENSE_ID", "UNLICENSED"),
        help="License identifier to embed in metadata."
    )
    parser.add_argument(
        "--build-id",
        default=os.getenv("BUILD_ID", ""),
        help="Build identifier to embed in metadata."
    )
    parser.add_argument("--workers", type=int, default=1, help="Number of worker threads.")
    parser.add_argument(
        "--output-layout",
        choices=["flat", "by-type"],
        default="flat",
        help="Output layout: flat (single directory) or by-type (subfolders per content type)."
    )
    parser.add_argument(
        "--output-mode",
        choices=["single", "role-split"],
        default="single",
        help="Output mode: single JSON or role-split (master/technical/marketing)."
    )
    parser.add_argument(
        "--output-format",
        choices=["json", "md", "both"],
        default="json",
        help="Output format: json, md, or both."
    )
    parser.add_argument("--log-level", default="INFO", help="Logging level.")
    return parser.parse_args()

def main() -> int:
    args = parse_args()
    logging.basicConfig(
        level=getattr(logging, args.log_level.upper(), logging.INFO),
        format="%(asctime)s | %(levelname)s | %(message)s"
    )

    input_dir = Path(args.input_dir)
    output_dir = Path(args.output_dir)
    output_dir.mkdir(exist_ok=True, parents=True)

    build_id = args.build_id or str(uuid.uuid4())
    client = configure_gemini(args.enable_llm)
    run_id = str(uuid.uuid4())
    workers = max(1, int(args.workers))

    if not input_dir.exists():
        logging.error(f"Input directory not found: {input_dir}")
        return 1

    files = sorted([p for p in input_dir.iterdir() if p.suffix.lower() in (".txt", ".md", ".pdf", ".docx")])
    if args.max_files and args.max_files > 0:
        files = files[:args.max_files]

    if not files:
        logging.info("No raw files found.")
        return 0

    success, fail = 0, 0
    results: List[Dict[str, object]] = []
    audit_entries: List[Dict[str, object]] = []
    output_lock = None
    if workers > 1:
        import threading
        output_lock = threading.Lock()

    run_started = datetime.datetime.utcnow()

    def handle_result(result: Dict[str, object], ok: bool, file_name: str) -> None:
        nonlocal success, fail
        results.append(result)
        status = result.get("status", "")
        if status not in ("ok", "dry_run"):
            logging.warning(f"Skipped {file_name}: {status}")
        success += 1 if ok else 0
        fail += 0 if ok else 1
        audit_entries.append({
            "run_id": run_id,
            "source_id": hash_identifier(str(result.get("source_file", ""))),
            "source_hash": result.get("source_hash", ""),
            "status": status,
            "timestamp": datetime.datetime.utcnow().isoformat() + "Z"
        })

    if workers > 1:
        from concurrent.futures import ThreadPoolExecutor, as_completed

        with ThreadPoolExecutor(max_workers=workers) as executor:
            future_map = {
                executor.submit(
                    process_one,
                    p,
                    output_dir,
                    client,
                    run_id,
                    args.redact_pii,
                    args.anonymize_source,
                    args.dry_run,
                    args.skip_unchanged,
                    args.max_bytes,
                    output_lock,
                    args.output_layout,
                    args.output_mode,
                    args.output_format,
                    args.license_id,
                    build_id
                ): p
                for p in files
            }
            for future in as_completed(future_map):
                p = future_map[future]
                try:
                    result, ok = future.result()
                except Exception as e:
                    result = {
                        "source_file": "",
                        "source_name": p.name,
                        "source_hash": "",
                        "status": "error",
                        "output_file": "",
                        "warnings": [],
                        "errors": [f"exception:{e}"],
                        "duration_ms": 0
                    }
                    ok = False
                handle_result(result, ok, p.name)
    else:
        for p in files:
            result, ok = process_one(
                p,
                output_dir,
                client,
                run_id,
                args.redact_pii,
                args.anonymize_source,
                args.dry_run,
                args.skip_unchanged,
                args.max_bytes,
                output_lock,
                args.output_layout,
                args.output_mode,
                args.output_format,
                args.license_id,
                build_id
            )
            handle_result(result, ok, p.name)

    results.sort(key=lambda r: r.get("source_name", ""))
    skipped = sum(1 for r in results if r.get("status") == "skipped")
    dry_runs = sum(1 for r in results if r.get("status") == "dry_run")
    warning_count = sum(len(r.get("warnings", [])) for r in results)
    error_count = sum(len(r.get("errors", [])) for r in results)

    run_finished = datetime.datetime.utcnow()
    summary = {
        "run_id": run_id,
        "pipeline": PIPELINE_NAME,
        "pipeline_version": PIPELINE_VERSION,
        "schema_version": SCHEMA_VERSION,
        "schema_path": Path(args.schema).as_posix(),
        "input_dir": input_dir.as_posix(),
        "output_dir": output_dir.as_posix(),
        "started_at": run_started.isoformat() + "Z",
        "finished_at": run_finished.isoformat() + "Z",
        "total_files": len(files),
        "success": success,
        "failed": fail,
        "skipped": skipped,
        "dry_run": dry_runs,
        "warnings": warning_count,
        "errors": error_count,
        "results": results
    }

    manifest_path = Path(args.manifest) if args.manifest else output_dir / "_manifest.json"
    if not args.dry_run:
        write_manifest(manifest_path, summary)
        logging.info(f"Manifest written: {manifest_path}")

    if args.audit_log:
        write_audit_log(Path(args.audit_log), audit_entries)

    logging.info(f"Done. Success: {success} | Failed: {fail}")
    return 0 if fail == 0 else 1

if __name__ == "__main__":
    raise SystemExit(main())
